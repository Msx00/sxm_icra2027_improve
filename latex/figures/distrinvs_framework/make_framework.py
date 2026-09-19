#!/usr/bin/env python3
"""Generate the DistriNVS framework figure and a fully editable PPTX.

The paper assets are drawn with Matplotlib (SVG/PDF/PNG/TIFF).  The PPTX is
rebuilt with native PowerPoint shapes and text, rather than embedding a raster
render, so that every module, arrow, label, and color remains editable.
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import rcParams
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle
from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# Visual language shared with the first-page teaser.
INK = "#172331"
MUTED = "#5C6877"
RULE = "#CBD4DE"
PANEL = "#F8FAFC"
TRAIN = "#F3F0F8"
BLUE = "#2F6FAE"
BLUE_LIGHT = "#E7F1FA"
PURPLE = "#7455A4"
PURPLE_LIGHT = "#F0EAF7"
TEAL = "#12857D"
TEAL_LIGHT = "#E3F3F0"
AMBER = "#D67723"
AMBER_LIGHT = "#FCEBD9"
TISSUE = "#E99B87"
TISSUE_LIGHT = "#F8DED5"
RED = "#B84D55"
WHITE = "#FFFFFF"

CANVAS_W = 180.0
CANVAS_H = 75.0
FIG_W_IN = 7.20  # 182.9 mm, IEEE double-column width
FIG_H_IN = 3.00  # 76.2 mm
SLIDE_W_IN = 13.333
SLIDE_H_IN = SLIDE_W_IN * CANVAS_H / CANVAS_W


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _mpl_setup() -> None:
    rcParams.update(
        {
            "font.family": "sans-serif",
            # DejaVu Sans covers the Greek symbols and Unicode subscripts used
            # in the compact variable labels.
            "font.sans-serif": ["DejaVu Sans", "Arial", "Liberation Sans"],
            "font.size": 6.0,
            "axes.linewidth": 0.6,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


class MplCanvas:
    """Small drawing adapter using top-left coordinates."""

    def __init__(self, ax):
        self.ax = ax

    def box(self, x, y, w, h, *, fill=WHITE, edge=RULE, lw=0.65, radius=1.2, z=2):
        patch = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle=f"round,pad=0.08,rounding_size={radius}",
            facecolor=fill,
            edgecolor=edge,
            linewidth=lw,
            zorder=z,
        )
        self.ax.add_patch(patch)
        return patch

    def rect(self, x, y, w, h, *, fill=WHITE, edge=RULE, lw=0.55, z=3):
        patch = Rectangle((x, y), w, h, facecolor=fill, edgecolor=edge, linewidth=lw, zorder=z)
        self.ax.add_patch(patch)
        return patch

    def ellipse(self, cx, cy, w, h, *, fill=WHITE, edge=RULE, lw=0.6, z=4, alpha=1.0):
        patch = Ellipse((cx, cy), w, h, facecolor=fill, edgecolor=edge, linewidth=lw, zorder=z, alpha=alpha)
        self.ax.add_patch(patch)
        return patch

    def polygon(self, points, *, fill=WHITE, edge=RULE, lw=0.55, z=3, alpha=1.0):
        patch = Polygon(points, closed=True, facecolor=fill, edgecolor=edge, linewidth=lw, zorder=z, alpha=alpha)
        self.ax.add_patch(patch)
        return patch

    def text(
        self,
        x,
        y,
        w,
        h,
        value,
        *,
        size=6.0,
        color=INK,
        bold=False,
        align="center",
        valign="center",
        italic=False,
        z=8,
    ):
        ha = {"left": "left", "center": "center", "right": "right"}[align]
        va = {"top": "top", "center": "center", "bottom": "bottom"}[valign]
        tx = x if align == "left" else x + w if align == "right" else x + w / 2
        ty = y if valign == "top" else y + h if valign == "bottom" else y + h / 2
        return self.ax.text(
            tx,
            ty,
            value,
            ha=ha,
            va=va,
            fontsize=size,
            color=color,
            weight="bold" if bold else "normal",
            style="italic" if italic else "normal",
            linespacing=1.05,
            zorder=z,
        )

    def line(self, x1, y1, x2, y2, *, color=INK, lw=0.75, dashed=False, z=5):
        kwargs = {"color": color, "lw": lw, "zorder": z, "solid_capstyle": "round"}
        if dashed:
            kwargs["linestyle"] = (0, (3.0, 2.1))
        return self.ax.plot([x1, x2], [y1, y2], **kwargs)[0]

    def arrow(self, x1, y1, x2, y2, *, color=INK, lw=0.8, dashed=False, z=6, scale=6.0):
        style = (0, (3.0, 2.1)) if dashed else "solid"
        patch = FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=scale,
            color=color,
            linewidth=lw,
            linestyle=style,
            shrinkA=0,
            shrinkB=0,
            zorder=z,
        )
        self.ax.add_patch(patch)
        return patch


class PptCanvas:
    """PowerPoint adapter; all content is constructed from native shapes."""

    FONT_SCALE = 1.62

    def __init__(self, slide):
        self.slide = slide

    @staticmethod
    def sx(x):
        return Inches(SLIDE_W_IN * x / CANVAS_W)

    @staticmethod
    def sy(y):
        return Inches(SLIDE_H_IN * y / CANVAS_H)

    @staticmethod
    def _style(shape, fill, edge, lw):
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(edge)
        shape.line.width = Pt(lw)
        return shape

    def box(self, x, y, w, h, *, fill=WHITE, edge=RULE, lw=0.65, radius=1.2, z=2):
        del radius, z
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, self.sx(x), self.sy(y), self.sx(w), self.sy(h)
        )
        shape.adjustments[0] = 0.08
        return self._style(shape, fill, edge, max(0.5, lw * 1.3))

    def rect(self, x, y, w, h, *, fill=WHITE, edge=RULE, lw=0.55, z=3):
        del z
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, self.sx(x), self.sy(y), self.sx(w), self.sy(h))
        return self._style(shape, fill, edge, max(0.45, lw * 1.3))

    def ellipse(self, cx, cy, w, h, *, fill=WHITE, edge=RULE, lw=0.6, z=4, alpha=1.0):
        del z, alpha
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.OVAL, self.sx(cx - w / 2), self.sy(cy - h / 2), self.sx(w), self.sy(h)
        )
        return self._style(shape, fill, edge, max(0.45, lw * 1.3))

    def polygon(self, points, *, fill=WHITE, edge=RULE, lw=0.55, z=3, alpha=1.0):
        del z, alpha
        xs = [p[0] for p in points]
        ys = [p[1] for p in points]
        # Native freeform support is less portable than a close equivalent.
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.FREEFORM if hasattr(MSO_SHAPE, "FREEFORM") else MSO_SHAPE.ISOSCELES_TRIANGLE,
            self.sx(min(xs)),
            self.sy(min(ys)),
            self.sx(max(xs) - min(xs)),
            self.sy(max(ys) - min(ys)),
        )
        return self._style(shape, fill, edge, max(0.45, lw * 1.3))

    def text(
        self,
        x,
        y,
        w,
        h,
        value,
        *,
        size=6.0,
        color=INK,
        bold=False,
        align="center",
        valign="center",
        italic=False,
        z=8,
    ):
        del z
        box = self.slide.shapes.add_textbox(self.sx(x), self.sy(y), self.sx(w), self.sy(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Pt(0.3)
        tf.margin_right = Pt(0.3)
        tf.margin_top = Pt(0)
        tf.margin_bottom = Pt(0)
        tf.vertical_anchor = {
            "top": MSO_ANCHOR.TOP,
            "center": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }[valign]
        p = tf.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = value
        run.font.name = "Arial"
        run.font.size = Pt(size * self.FONT_SCALE)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
        return box

    def line(self, x1, y1, x2, y2, *, color=INK, lw=0.75, dashed=False, z=5):
        del z
        line = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, self.sx(x1), self.sy(y1), self.sx(x2), self.sy(y2)
        )
        line.line.color.rgb = rgb(color)
        line.line.width = Pt(max(0.55, lw * 1.35))
        if dashed:
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return line

    def arrow(self, x1, y1, x2, y2, *, color=INK, lw=0.8, dashed=False, z=6, scale=6.0):
        del z
        line = self.line(x1, y1, x2, y2, color=color, lw=lw, dashed=dashed)
        # python-pptx has no fully portable arrowhead API; a native triangle is
        # used as a separately editable arrowhead.
        angle = math.degrees(math.atan2(y2 - y1, x2 - x1))
        head_w = max(1.15, scale * 0.22)
        head_h = max(1.15, scale * 0.22)
        head = self.slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE,
            self.sx(x2 - head_w / 2),
            self.sy(y2 - head_h / 2),
            self.sx(head_w),
            self.sy(head_h),
        )
        self._style(head, color, color, 0.3)
        head.rotation = angle + 90.0
        return line, head


def panel(c, x, y, w, h, number, title, accent):
    c.box(x, y, w, h, fill=PANEL, edge=RULE, lw=0.65, radius=1.3, z=1)
    c.rect(x, y, w, 5.3, fill=WHITE, edge=WHITE, lw=0, z=2)
    c.line(x + 0.8, y + 5.3, x + w - 0.8, y + 5.3, color=RULE, lw=0.55, z=3)
    c.ellipse(x + 3.2, y + 2.65, 3.4, 3.4, fill=accent, edge=accent, lw=0.4, z=4)
    c.text(x + 1.55, y + 0.95, 3.3, 3.3, str(number), size=6.3, color=WHITE, bold=True)
    c.text(x + 5.4, y + 0.55, w - 6.2, 4.2, title, size=6.0, color=accent, bold=True, align="left")


def chip(c, x, y, w, text_value, *, fill=WHITE, edge=RULE, color=INK, size=5.2, bold=False):
    c.box(x, y, w, 4.4, fill=fill, edge=edge, lw=0.52, radius=1.0, z=4)
    c.text(x + 0.25, y + 0.2, w - 0.5, 4.0, text_value, size=size, color=color, bold=bold)


def draw_rgb_tile(c, x, y, w, h, label="RGB"):
    c.box(x, y, w, h, fill=TISSUE_LIGHT, edge=BLUE, lw=0.65, radius=0.7, z=3)
    c.line(x + 0.7, y + h * 0.67, x + w * 0.35, y + h * 0.45, color=TISSUE, lw=1.4, z=4)
    c.line(x + w * 0.35, y + h * 0.45, x + w * 0.62, y + h * 0.64, color=TISSUE, lw=1.4, z=4)
    c.line(x + w * 0.62, y + h * 0.64, x + w - 0.7, y + h * 0.38, color=TISSUE, lw=1.4, z=4)
    c.ellipse(x + w * 0.72, y + h * 0.34, 1.0, 1.0, fill=WHITE, edge=WHITE, lw=0.3, z=5)
    c.text(x + 0.2, y + h - 3.2, w - 0.4, 2.8, label, size=4.15, color=BLUE, bold=True)


def draw_depth_tile(c, x, y, w, h, label="DEPTH"):
    c.box(x, y, w, h, fill=BLUE_LIGHT, edge=BLUE, lw=0.65, radius=0.7, z=3)
    for i, col in enumerate(["#CADFF1", "#94BEDF", "#5E96C5", BLUE]):
        c.rect(x + 0.65 + i * (w - 1.3) / 4, y + 0.8, (w - 1.3) / 4 + 0.05, h - 3.6, fill=col, edge=col, lw=0.1, z=4)
    c.text(x + 0.2, y + h - 3.2, w - 0.4, 2.8, label, size=4.15, color=BLUE, bold=True)


def draw_mask_tile(c, x, y, w, h, label="MASK"):
    c.box(x, y, w, h, fill=WHITE, edge=BLUE, lw=0.65, radius=0.7, z=3)
    c.ellipse(x + w * 0.34, y + h * 0.36, 2.1, 2.1, fill=BLUE, edge=BLUE, lw=0.2, z=4)
    c.rect(x + w * 0.57, y + h * 0.22, 1.7, 3.2, fill=BLUE, edge=BLUE, lw=0.2, z=4)
    c.text(x + 0.2, y + h - 3.2, w - 0.4, 2.8, label, size=4.15, color=BLUE, bold=True)


def draw_stat_tile(c, x, y, w, h, label, color, kind):
    c.box(x, y, w, h, fill=WHITE, edge=RULE, lw=0.55, radius=0.65, z=3)
    area_y = y + 0.65
    area_h = h - 3.1
    if kind == "warp":
        c.rect(x + 0.55, area_y, w - 1.1, area_h, fill=TISSUE_LIGHT, edge=TISSUE_LIGHT, lw=0.2, z=4)
        c.line(x + 0.8, area_y + area_h * 0.7, x + w * 0.48, area_y + area_h * 0.35, color=TISSUE, lw=1.15)
        c.line(x + w * 0.48, area_y + area_h * 0.35, x + w - 0.8, area_y + area_h * 0.62, color=TISSUE, lw=1.15)
    elif kind == "depth":
        for i, col in enumerate(["#D6E7F4", "#A9CCE7", "#73A9D2", BLUE]):
            c.rect(x + 0.55 + i * (w - 1.1) / 4, area_y, (w - 1.1) / 4 + 0.05, area_h, fill=col, edge=col, lw=0.1)
    elif kind == "support":
        for i, hh in enumerate([0.9, 1.8, 2.7, 2.1]):
            c.rect(x + 1.0 + i * 1.75, area_y + area_h - hh, 1.05, hh, fill=TEAL, edge=TEAL, lw=0.15)
    elif kind == "variance":
        for i, rr in enumerate([0.7, 1.1, 1.55]):
            c.ellipse(x + 2.2 + i * 2.35, area_y + area_h / 2, rr * 1.2, rr, fill=AMBER_LIGHT, edge=AMBER, lw=0.45)
    elif kind == "collision":
        c.line(x + 0.9, area_y + area_h * 0.25, x + w - 0.9, area_y + area_h * 0.75, color=PURPLE, lw=1.0)
        c.line(x + 0.9, area_y + area_h * 0.75, x + w - 0.9, area_y + area_h * 0.25, color=PURPLE, lw=1.0)
        c.ellipse(x + w / 2, area_y + area_h / 2, 1.4, 1.4, fill=PURPLE, edge=WHITE, lw=0.3)
    c.text(x + 0.2, y + h - 2.35, w - 0.4, 1.8, label, size=4.55, color=color, bold=True)


def draw_framework(c) -> None:
    # Main four-stage structure.
    panel(c, 1.0, 1.0, 29.5, 54.8, 1, "SOURCE RELIABILITY", BLUE)
    panel(c, 31.5, 1.0, 60.0, 54.8, 2, "DISTRIBUTIONAL SURFACE SPLATTING", PURPLE)
    panel(c, 92.5, 1.0, 50.5, 54.8, 3, "VISIBILITY-GUIDED ROUTING", AMBER)
    panel(c, 144.0, 1.0, 35.0, 54.8, 4, "FUSED OUTPUT", TEAL)

    # Stage 1: source evidence and reliability-aware geometry.
    draw_rgb_tile(c, 3.1, 8.3, 7.4, 9.2, "RGB\nIₛ")
    draw_depth_tile(c, 12.0, 8.3, 7.4, 9.2, "SPARSE\nDₛ")
    draw_mask_tile(c, 20.9, 8.3, 7.4, 9.2, "VALID\nMₛ")
    c.box(3.1, 19.2, 25.2, 5.2, fill=WHITE, edge=BLUE, lw=0.55, radius=0.8)
    c.text(3.5, 19.55, 24.4, 4.5, "calibration  Kₛ · Kₜ · Tₛ→ₜ", size=5.2, color=BLUE, bold=True)

    c.arrow(15.7, 24.6, 15.7, 27.0, color=BLUE, lw=0.7)
    c.box(3.1, 27.2, 25.2, 10.7, fill=BLUE_LIGHT, edge=BLUE, lw=0.7, radius=1.0)
    c.text(4.1, 27.8, 23.2, 2.7, "DEPTH RELIABILITY ENCODER", size=4.55, color=BLUE, bold=True)
    c.box(4.5, 31.2, 9.5, 4.7, fill=WHITE, edge=RULE, lw=0.5, radius=0.7)
    c.text(4.8, 31.45, 8.9, 4.1, "full-res\nresidual", size=4.8, color=INK)
    c.box(17.3, 31.2, 9.5, 4.7, fill=WHITE, edge=RULE, lw=0.5, radius=0.7)
    c.text(17.6, 31.45, 8.9, 4.1, "¼-res\ncontext", size=4.8, color=INK)
    c.arrow(14.2, 33.55, 17.0, 33.55, color=BLUE, lw=0.55, scale=5.0)

    for x, w, txt, col in [
        (3.2, 4.6, "μ", BLUE),
        (8.2, 4.6, "σ", PURPLE),
        (13.2, 4.6, "n", TEAL),
        (18.2, 4.6, "c", AMBER),
        (23.2, 5.0, "Fₛ", INK),
    ]:
        chip(c, x, 40.2, w, txt, fill=WHITE, edge=col, color=col, size=5.6, bold=True)
    c.text(
        3.2,
        44.9,
        25.0,
        3.8,
        "μ mean · σ scale · n normal\nc confidence · Fₛ feature",
        size=4.05,
        color=MUTED,
    )
    # A compact distribution glyph reinforces that depth is not a point value.
    c.line(5.0, 52.2, 26.4, 52.2, color=RULE, lw=0.5)
    pts = []
    for i in range(28):
        xx = 6.0 + i * 0.70
        yy = 52.0 - 3.0 * math.exp(-0.5 * ((xx - 15.5) / 3.2) ** 2)
        pts.append((xx, yy))
    for p1, p2 in zip(pts[:-1], pts[1:]):
        c.line(p1[0], p1[1], p2[0], p2[1], color=PURPLE, lw=0.8)
    for xx, lab in [(12.2, "−σ"), (15.5, "μ"), (18.8, "+σ")]:
        c.ellipse(xx, 52.2, 1.25, 1.25, fill=PURPLE, edge=WHITE, lw=0.3)
        c.text(xx - 1.5, 53.0, 3.0, 1.8, lab, size=4.3, color=PURPLE, bold=True)

    # Stage-to-stage feature flow.
    c.arrow(28.8, 33.0, 33.0, 33.0, color=INK, lw=0.9, scale=6.5)

    # Stage 2: DSS hero schematic.
    c.box(34.1, 8.4, 15.7, 5.0, fill=PURPLE_LIGHT, edge=PURPLE, lw=0.6, radius=0.8)
    c.text(34.6, 8.7, 14.7, 4.4, "Aₛ = [Iₛ, Fₛ]", size=5.3, color=PURPLE, bold=True)

    # Source ray and deterministic sigma samples.
    c.text(34.3, 15.0, 18.5, 2.5, "DETERMINISTIC σ SAMPLES", size=4.45, color=MUTED, bold=True)
    c.ellipse(36.3, 24.1, 2.2, 2.2, fill=BLUE, edge=WHITE, lw=0.35)
    c.line(37.4, 24.1, 54.0, 24.1, color=BLUE, lw=0.8)
    c.polygon([(36.5, 23.0), (52.0, 19.1), (52.0, 29.1)], fill=BLUE_LIGHT, edge=BLUE_LIGHT, lw=0.2, z=2, alpha=0.7)
    for xx, rr, lab in [(44.0, 1.3, "−σ"), (48.0, 1.6, "μ"), (52.0, 1.3, "+σ")]:
        c.ellipse(xx, 24.1, rr, rr, fill=PURPLE, edge=WHITE, lw=0.3)
        c.text(xx - 1.8, 25.0, 3.6, 1.8, lab, size=4.2, color=PURPLE, bold=True)
    c.text(34.7, 27.9, 17.5, 2.3, "zqk = clip(μ + ξkσ)", size=4.55, color=PURPLE)

    c.arrow(54.0, 24.1, 59.0, 24.1, color=PURPLE, lw=0.8)
    c.box(58.9, 16.0, 11.8, 16.5, fill=WHITE, edge=PURPLE, lw=0.65, radius=0.9)
    c.text(59.6, 16.6, 10.4, 3.1, "CALIBRATED\nPROJECTION", size=4.35, color=PURPLE, bold=True)
    c.text(59.7, 21.0, 10.2, 4.2, "Kₛ → Tₛ→ₜ → Kₜ", size=4.9, color=INK)
    c.line(61.0, 27.4, 68.4, 22.3, color=PURPLE, lw=0.7)
    c.line(61.0, 27.4, 68.4, 25.8, color=PURPLE, lw=0.7)
    c.line(61.0, 27.4, 68.4, 29.3, color=PURPLE, lw=0.7)
    c.text(59.7, 29.0, 10.2, 2.3, "Jacobian + tilt", size=4.45, color=MUTED)

    # Target-plane anisotropic footprints.
    c.arrow(70.8, 24.1, 73.5, 24.1, color=PURPLE, lw=0.8)
    c.box(73.6, 15.2, 15.0, 18.0, fill=WHITE, edge=PURPLE, lw=0.65, radius=0.9)
    c.text(74.1, 15.4, 14.0, 3.4, "TARGET-PLANE\nSPLATS", size=4.25, color=PURPLE, bold=True)
    c.rect(75.6, 19.0, 11.0, 9.3, fill=PURPLE_LIGHT, edge=RULE, lw=0.45)
    for cx, cy, ww, hh, col in [
        (78.1, 22.6, 5.8, 2.0, BLUE),
        (81.0, 24.7, 7.0, 2.5, PURPLE),
        (83.7, 21.3, 5.2, 1.8, TEAL),
    ]:
        c.ellipse(cx, cy, ww, hh, fill=WHITE, edge=col, lw=0.75, alpha=0.75)
    c.text(74.3, 28.8, 13.7, 3.4, "anisotropic footprint\nsoft nearest surface", size=4.45, color=MUTED)

    # DSS output maps: geometry and uncertainty share the same rasterization.
    c.text(
        34.2,
        33.9,
        54.5,
        3.4,
        "ONE RASTERIZATION RETURNS BOTH\nwarped evidence + visibility statistics",
        size=4.55,
        color=PURPLE,
        bold=True,
    )
    tile_y, tile_w, tile_h = 38.0, 10.3, 8.8
    for i, (lab, col, kind) in enumerate(
        [
            ("warp  Iᵂ,Fᵂ", TEAL, "warp"),
            ("depth  D̄", BLUE, "depth"),
            ("support  S,C", TEAL, "support"),
            ("variance  V", AMBER, "variance"),
            ("collision  H", PURPLE, "collision"),
        ]
    ):
        draw_stat_tile(c, 34.2 + i * 11.0, tile_y, tile_w, tile_h, lab, col, kind)
    c.box(45.2, 49.1, 32.0, 4.7, fill=PURPLE_LIGHT, edge=PURPLE, lw=0.55, radius=0.8)
    c.text(45.7, 49.45, 31.0, 4.0, "target reliability  κ = f(C, V, H)", size=5.0, color=PURPLE, bold=True)

    # Stage 3: reliability-guided route and two experts.
    stat_labels = [("C", TEAL), ("κ", PURPLE), ("V", AMBER), ("H", PURPLE), ("D̄", BLUE), ("M", MUTED)]
    for i, (lab, col) in enumerate(stat_labels):
        chip(c, 95.2 + i * 7.35, 8.3, 6.1, lab, fill=WHITE, edge=col, color=col, size=5.4, bold=True)
    c.text(96.0, 12.8, 43.9, 1.7, "renderer-derived condition h", size=4.55, color=MUTED)

    # Warped evidence is a data input to the transport expert; the gate only
    # determines how much of that estimate reaches the constrained mixture.
    c.line(91.6, 16.2, 118.0, 16.2, color=TEAL, lw=0.7, z=5)
    c.arrow(118.0, 16.2, 119.3, 17.0, color=TEAL, lw=0.7, scale=5.0)
    c.text(97.0, 14.65, 13.5, 1.4, "warped evidence", size=4.15, color=TEAL, bold=True)

    # The statistics themselves feed the routing condition.
    c.line(89.5, 42.2, 93.8, 42.2, color=PURPLE, lw=0.7)
    c.line(93.8, 42.2, 93.8, 10.5, color=PURPLE, lw=0.7)
    c.arrow(93.8, 10.5, 94.9, 10.5, color=PURPLE, lw=0.7, scale=5.0)

    c.arrow(112.8, 15.2, 112.8, 18.0, color=AMBER, lw=0.75)
    c.box(96.0, 18.2, 20.2, 18.2, fill=AMBER_LIGHT, edge=AMBER, lw=0.7, radius=1.0)
    c.text(97.0, 18.9, 18.2, 3.2, "BOUNDED ROUTER", size=5.7, color=AMBER, bold=True)
    c.box(98.0, 23.1, 7.2, 6.3, fill=WHITE, edge=AMBER, lw=0.5, radius=0.7)
    c.text(98.3, 23.4, 6.6, 5.7, "physics\nprior g⁰", size=4.7, color=INK)
    c.text(105.5, 24.3, 2.1, 3.8, "+", size=6.2, color=AMBER, bold=True)
    c.box(108.0, 23.1, 6.2, 6.3, fill=WHITE, edge=AMBER, lw=0.5, radius=0.7)
    c.text(108.3, 23.4, 5.6, 5.7, "learned\nΔg", size=4.7, color=INK)
    c.box(99.0, 31.0, 14.2, 3.4, fill=AMBER, edge=AMBER, lw=0.3, radius=1.2)
    c.text(99.2, 31.15, 13.8, 3.0, "g ∈ [0,1]", size=5.0, color=WHITE, bold=True)

    c.arrow(116.5, 25.5, 119.3, 22.7, color=TEAL, lw=0.75)
    c.arrow(116.5, 29.2, 119.3, 35.2, color=AMBER, lw=0.75)
    c.text(116.0, 20.2, 4.0, 2.0, "1−g", size=4.5, color=TEAL, bold=True)
    c.text(116.6, 35.0, 2.6, 2.0, "g", size=4.5, color=AMBER, bold=True)

    c.box(119.5, 17.0, 20.5, 11.6, fill=TEAL_LIGHT, edge=TEAL, lw=0.7, radius=1.0)
    c.text(120.4, 17.6, 18.7, 3.0, "TRANSPORT EXPERT", size=5.05, color=TEAL, bold=True)
    c.text(120.4, 21.0, 18.7, 4.8, "warped evidence\n+ bounded ΔRGB", size=4.9, color=INK)

    c.box(119.5, 31.0, 20.5, 12.9, fill=AMBER_LIGHT, edge=AMBER, lw=0.7, radius=1.0)
    c.text(120.4, 31.6, 18.7, 3.0, "SYNTHESIS EXPERT", size=5.05, color=AMBER, bold=True)
    c.text(120.4, 35.0, 18.7, 5.8, "local conv + FFT\nRGB · depth · risk", size=4.9, color=INK)

    c.line(140.2, 22.8, 141.8, 22.8, color=TEAL, lw=0.7)
    c.line(141.8, 22.8, 141.8, 49.8, color=TEAL, lw=0.7)
    c.arrow(141.8, 49.8, 140.8, 49.8, color=TEAL, lw=0.7, scale=5.0)
    c.arrow(129.7, 44.2, 129.7, 46.1, color=AMBER, lw=0.7)
    c.box(118.0, 46.3, 22.8, 7.0, fill=WHITE, edge=INK, lw=0.7, radius=1.0)
    c.text(118.8, 46.8, 21.2, 2.4, "SOFT FUSION", size=5.1, color=INK, bold=True)
    c.text(119.0, 49.2, 20.8, 3.1, "(1−g) transport + g synthesis", size=4.3, color=MUTED)

    # The reliability-constrained mixture is the final learned RGB-D prediction.
    c.arrow(141.0, 49.8, 147.0, 38.2, color=INK, lw=0.9, scale=6.5)

    # Stage 4: direct fused prediction; no output hard copy.
    c.box(147.0, 28.0, 29.0, 12.4, fill=WHITE, edge=INK, lw=0.75, radius=1.0)
    c.text(148.0, 28.6, 27.0, 3.0, "FUSED OUTPUT", size=5.6, color=INK, bold=True)
    c.text(148.0, 32.0, 27.0, 5.5, "blended RGB-D\n+ predicted risk", size=4.8, color=MUTED)

    c.box(148.0, 43.0, 27.0, 9.9, fill=WHITE, edge=TEAL, lw=0.75, radius=0.9)
    c.rect(149.0, 44.0, 11.8, 5.2, fill=TEAL_LIGHT, edge=TEAL_LIGHT, lw=0.2)
    c.rect(160.8, 44.0, 13.2, 5.2, fill=AMBER_LIGHT, edge=AMBER_LIGHT, lw=0.2)
    c.line(149.8, 47.4, 173.0, 45.7, color=TISSUE, lw=1.3)
    for i, hh in enumerate([1.0, 1.7, 2.4]):
        c.rect(166.5 + i * 1.45, 48.3 - hh, 0.75, hh, fill=AMBER, edge=AMBER, lw=0.15)
    c.text(148.2, 49.8, 26.6, 2.5, "Îₜ  ·  D̂ₜ  ·  Uₜ", size=5.25, color=INK, bold=True)
    c.arrow(161.5, 40.7, 161.5, 42.6, color=INK, lw=0.8)

    # Training-only lane.
    c.box(1.0, 58.0, 178.0, 15.8, fill=TRAIN, edge=PURPLE, lw=0.65, radius=1.2, z=1)
    c.box(3.0, 59.5, 25.0, 4.2, fill=PURPLE, edge=PURPLE, lw=0.3, radius=1.1, z=3)
    c.text(3.4, 59.75, 24.2, 3.6, "TRAINING ONLY · DASHED", size=4.55, color=WHITE, bold=True)

    c.box(29.0, 64.8, 28.0, 6.2, fill=WHITE, edge=PURPLE, lw=0.55, radius=0.8)
    c.text(29.7, 65.15, 26.6, 2.4, "target RGB-D + mask", size=5.0, color=PURPLE, bold=True)
    c.text(29.7, 67.7, 26.6, 2.4, "supervision; not an inference input", size=4.35, color=MUTED)

    c.box(62.0, 64.8, 37.0, 6.2, fill=WHITE, edge=PURPLE, lw=0.55, radius=0.8)
    c.text(62.7, 65.2, 35.6, 5.5, "Lrgb · Ldepth · Lrouter · Lrisk", size=4.9, color=PURPLE, bold=True)

    c.box(105.0, 64.8, 27.0, 6.2, fill=WHITE, edge=PURPLE, lw=0.55, radius=0.8)
    c.text(105.7, 65.2, 25.6, 5.5, "BACKWARD DSS\nTₜ→ₛ", size=4.85, color=PURPLE, bold=True)

    c.box(138.0, 64.8, 34.0, 6.2, fill=WHITE, edge=PURPLE, lw=0.55, radius=0.8)
    c.text(138.7, 65.2, 32.6, 5.5, "source-view cycle  Lcyc", size=4.55, color=PURPLE, bold=True)

    c.arrow(57.4, 67.9, 61.5, 67.9, color=PURPLE, lw=0.65, dashed=True, scale=5.5)
    c.arrow(132.4, 67.9, 137.5, 67.9, color=PURPLE, lw=0.65, dashed=True, scale=5.5)
    c.arrow(137.6, 70.9, 99.4, 70.9, color=PURPLE, lw=0.65, dashed=True, scale=5.5)
    c.arrow(161.5, 53.2, 119.0, 64.4, color=PURPLE, lw=0.7, dashed=True, scale=5.5)
    # Mask the diagonal behind its label so the text stays readable in both
    # PowerPoint and LibreOffice rendering.
    c.box(111.7, 58.7, 32.6, 3.9, fill=TRAIN, edge=TRAIN, lw=0.0, radius=0.2)
    c.text(112.0, 58.9, 32.0, 3.5, "predicted RGB-D + risk", size=4.55, color=PURPLE, bold=True)


def draw_mpl(out_dir: Path) -> None:
    _mpl_setup()
    fig = plt.figure(figsize=(FIG_W_IN, FIG_H_IN), facecolor=WHITE)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, CANVAS_W)
    ax.set_ylim(CANVAS_H, 0)
    ax.axis("off")
    canvas = MplCanvas(ax)
    draw_framework(canvas)

    out_dir.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_dir / "distrinvs_framework.svg", format="svg", facecolor=WHITE)
    fig.savefig(out_dir / "distrinvs_framework.pdf", format="pdf", facecolor=WHITE)
    fig.savefig(out_dir / "distrinvs_framework.png", format="png", dpi=600, facecolor=WHITE)
    fig.savefig(
        out_dir / "distrinvs_framework.tiff",
        format="tiff",
        dpi=600,
        facecolor=WHITE,
        pil_kwargs={"compression": "tiff_lzw"},
    )
    plt.close(fig)


def add_editing_guide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    c = PptCanvas(slide)
    c.text(5, 5, 170, 7, "DistriNVS framework · editing guide", size=14, color=INK, bold=True, align="left")
    c.text(
        5,
        13,
        170,
        10,
        "Recommended paper size: 183 × 76 mm (double column). Keep the renderer-derived statistics, constrained gate, two experts, fused output, and training-only cycle visually distinct.",
        size=8.3,
        color=MUTED,
        align="left",
    )
    palette = [
        (BLUE, "source / depth"),
        (PURPLE, "DSS / geometry"),
        (TEAL, "transport / fusion"),
        (AMBER, "synthesis / risk"),
        (TISSUE, "endoscopic appearance"),
    ]
    for i, (col, name) in enumerate(palette):
        x = 6 + i * 34.2
        c.box(x, 29, 27, 9, fill=col, edge=col, lw=0.4, radius=1.2)
        c.text(x - 2, 39, 31, 5, name, size=7.0, color=INK, bold=True)
    c.text(5, 49, 170, 5, "Suggested caption", size=9.5, color=INK, bold=True, align="left")
    c.text(
        5,
        55,
        170,
        15,
        "DistriNVS framework. A reliability encoder predicts source geometry and features. DSS transports deterministic depth hypotheses through the calibrated cross-endoscope transform and jointly produces warped evidence and target-plane visibility statistics. These statistics condition a reliability-constrained soft mixture of a residual transport expert and a Fourier synthesis expert. Solid arrows denote inference; the dashed lower lane is used only during training.",
        size=7.4,
        color=MUTED,
        align="left",
    )


def add_picture_tile(
    c: PptCanvas,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    accent: str,
    *,
    dashed: bool = False,
) -> None:
    """Place a 5:4 evidence image without cropping its invalid black region."""
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as image:
        image_w, image_h = image.size
    if image_w <= 0 or image_h <= 0:
        raise ValueError(f"invalid image dimensions: {path}")
    target_ratio = w / h
    image_ratio = image_w / image_h
    if abs(target_ratio - image_ratio) > 1.0e-3:
        raise ValueError(
            f"evidence tile would distort {path.name}: "
            f"source={image_ratio:.5f}, target={target_ratio:.5f}"
        )
    picture = c.slide.shapes.add_picture(
        str(path), c.sx(x), c.sy(y), width=c.sx(w), height=c.sy(h)
    )
    picture.name = f"Evidence · {label}"
    frame = c.slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, c.sx(x), c.sy(y), c.sx(w), c.sy(h)
    )
    frame.fill.background()
    frame.line.color.rgb = rgb(accent)
    frame.line.width = Pt(0.8)
    if dashed:
        frame.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    c.text(x, y + h + 0.35, w, 3.0, label, size=4.35, color=accent, bold=True)


def add_nohard_evidence_slide(prs: Presentation, assets: Path) -> None:
    """Add a traceable real-case walkthrough without crowding the framework."""
    required = {
        "requested": "00_requested_tool4_train_exemplar.png",
        "source": "01_matched_source_rgb.png",
        "warp": "02_dss_warp.png",
        "support": "05_dss_support.png",
        "variance": "06_dss_variance.png",
        "collision": "07_collision_entropy.png",
        "prior": "08_physics_prior.png",
        "transport": "09_transport_rgb.png",
        "synthesis": "10_synthesis_rgb.png",
        "gate": "11_synthesis_gate.png",
        "fused": "12_soft_fused_rgb.png",
        "risk": "13_predicted_risk.png",
    }
    paths = {key: assets / filename for key, filename in required.items()}
    missing = [str(path) for path in paths.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"missing no-hard evidence assets: {missing}")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    c = PptCanvas(slide)
    c.text(
        3.0,
        1.1,
        174.0,
        5.2,
        "NO-HARD COMPOSITION · TRACEABLE SINGLE-FRAME INTERNALS",
        size=10.2,
        color=INK,
        bold=True,
        align="left",
    )
    c.text(
        3.0,
        6.0,
        174.0,
        3.0,
        "Matched forward pass: tool_3 / frame_000002 · step 10,000 · no post-processing · tool_4 is shown separately as a train-split appearance example",
        size=4.25,
        color=MUTED,
        align="left",
    )

    panel(c, 1.0, 10.0, 32.0, 62.5, 1, "INPUT CONTEXT", BLUE)
    panel(c, 34.0, 10.0, 56.0, 62.5, 2, "DSS GEOMETRY", PURPLE)
    panel(c, 91.0, 10.0, 58.0, 62.5, 3, "SOFT ROUTING + EXPERTS", AMBER)
    panel(c, 150.0, 10.0, 29.0, 62.5, 4, "OUTPUT", TEAL)

    # The user-selected tool_4 RGB is retained, but is explicitly separated
    # from the matched held-out tool_3 inference chain.
    add_picture_tile(
        c,
        paths["source"],
        3.5,
        17.0,
        27.0,
        21.6,
        "matched tool_3 source  Iₛ",
        BLUE,
    )
    add_picture_tile(
        c,
        paths["requested"],
        3.5,
        44.0,
        27.0,
        21.6,
        "tool_4 RGB · TRAIN EXAMPLE",
        BLUE,
        dashed=True,
    )

    for path, x, y, label in (
        (paths["warp"], 36.5, 17.0, "DSS warp  Iₜʷ"),
        (paths["support"], 63.5, 17.0, "support  S"),
        (paths["variance"], 36.5, 44.0, "variance  V"),
        (paths["collision"], 63.5, 44.0, "collision entropy  H"),
    ):
        add_picture_tile(c, path, x, y, 24.0, 19.2, label, PURPLE)

    for path, x, y, label, accent in (
        (paths["prior"], 93.5, 17.0, "physics prior  g⁰", PURPLE),
        (paths["gate"], 121.5, 17.0, "synthesis gate  g", AMBER),
        (paths["transport"], 93.5, 44.0, "transport expert", TEAL),
        (paths["synthesis"], 121.5, 44.0, "synthesis expert", AMBER),
    ):
        add_picture_tile(c, path, x, y, 25.0, 20.0, label, accent)

    add_picture_tile(
        c,
        paths["fused"],
        152.5,
        17.0,
        24.0,
        19.2,
        "soft-fused RGB",
        TEAL,
    )
    add_picture_tile(
        c,
        paths["risk"],
        152.5,
        44.0,
        24.0,
        19.2,
        "predicted risk  Uₜ",
        AMBER,
    )

    # Small directional cues reinforce the four-stage reading order without
    # crossing the images or their labels.
    c.arrow(32.2, 41.0, 33.8, 41.0, color=INK, lw=0.75, scale=5.5)
    c.arrow(89.2, 41.0, 90.8, 41.0, color=INK, lw=0.75, scale=5.5)
    c.arrow(148.2, 41.0, 149.8, 41.0, color=INK, lw=0.75, scale=5.5)


def draw_pptx(out_dir: Path, evidence_dir: Path | None = None) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    draw_framework(PptCanvas(slide))
    if evidence_dir is not None:
        add_nohard_evidence_slide(prs, evidence_dir)
    add_editing_guide(prs)
    prs.save(out_dir / "distrinvs_framework_editable.pptx")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument(
        "--evidence-dir",
        type=Path,
        default=None,
        help="directory created by extract_nohard_intermediates.py",
    )
    parser.add_argument(
        "--pptx-only",
        action="store_true",
        help="rebuild only the editable deck and leave paper figure exports untouched",
    )
    args = parser.parse_args()
    output_dir = args.output_dir.expanduser().resolve()
    evidence_dir = (
        args.evidence_dir.expanduser().resolve()
        if args.evidence_dir is not None
        else output_dir / "intermediate_assets"
    )
    if not args.pptx_only:
        draw_mpl(output_dir)
    draw_pptx(output_dir, evidence_dir if evidence_dir.is_dir() else None)
    print(f"Wrote framework assets to {args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
