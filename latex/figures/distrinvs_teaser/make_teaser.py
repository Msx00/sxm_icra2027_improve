#!/usr/bin/env python3
"""Create the DistriNVS first-page teaser and an editable PowerPoint version."""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import rcParams
from matplotlib.patches import (
    Circle,
    Ellipse,
    FancyArrowPatch,
    FancyBboxPatch,
    PathPatch,
    Polygon,
    Rectangle,
    Wedge,
)
from matplotlib.path import Path as MplPath
from matplotlib.transforms import Affine2D

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


INK = "#172331"
MUTED = "#5C6877"
RULE = "#CBD4DE"
BLUE = "#2F6FAE"
BLUE_LIGHT = "#E7F1FA"
PURPLE = "#7455A4"
PURPLE_LIGHT = "#F0EAF7"
TEAL = "#12857D"
TEAL_LIGHT = "#E3F3F0"
AMBER = "#D67723"
AMBER_LIGHT = "#FCEBD9"
TISSUE = "#E99B87"
TISSUE_LIGHT = "#F8DED5"
WHITE = "#FFFFFF"

CANVAS_W = 100.0
CANVAS_H = 66.7
FIG_W_IN = 3.48  # 88.4 mm: IEEE single-column width
FIG_H_IN = 2.32


def _mpl_setup() -> None:
    rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Liberation Sans", "DejaVu Sans"],
            "font.size": 7.0,
            "axes.linewidth": 0.6,
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


def rounded_box(ax, xy, width, height, *, face, edge, radius=1.8, lw=0.8, z=2):
    patch = FancyBboxPatch(
        xy,
        width,
        height,
        boxstyle=f"round,pad=0.15,rounding_size={radius}",
        facecolor=face,
        edgecolor=edge,
        linewidth=lw,
        zorder=z,
    )
    ax.add_patch(patch)
    return patch


def draw_mpl_endoscope(ax, x, y, angle_deg, color, label):
    transform = Affine2D().rotate_deg_around(x, y, angle_deg) + ax.transData
    body = FancyBboxPatch(
        (x - 6.5, y - 2.6),
        12.5,
        5.2,
        boxstyle="round,pad=0.1,rounding_size=1.6",
        facecolor=WHITE,
        edgecolor=color,
        linewidth=1.2,
        transform=transform,
        zorder=6,
    )
    ax.add_patch(body)
    ax.add_patch(
        Rectangle(
            (x - 4.0, y - 1.2),
            4.2,
            2.4,
            facecolor=color,
            edgecolor="none",
            transform=transform,
            zorder=7,
        )
    )
    ax.plot([x + 5.5, x + 14.0], [y, y], color=color, lw=2.1, transform=transform, zorder=6)
    ax.add_patch(
        Ellipse(
            (x + 14.1, y),
            2.0,
            3.2,
            facecolor=INK,
            edgecolor=WHITE,
            linewidth=0.5,
            transform=transform,
            zorder=7,
        )
    )
    theta = math.radians(angle_deg)
    lens = (x + 14.1 * math.cos(theta), y + 14.1 * math.sin(theta))
    ax.text(x, y + 5.0, label, ha="center", va="bottom", color=color, fontsize=6.0, weight="bold")
    return lens


def tissue_path():
    verts = [
        (5, 34.1),
        (14, 37.2),
        (23, 36.0),
        (31, 39.0),
        (42, 36.8),
        (52, 39.5),
        (62, 37.4),
        (73, 40.2),
        (84, 37.8),
        (95, 39.2),
        (98, 37.0),
        (98, 32.2),
        (5, 32.2),
        (5, 34.1),
    ]
    codes = [
        MplPath.MOVETO,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.CURVE3,
        MplPath.LINETO,
        MplPath.LINETO,
        MplPath.LINETO,
        MplPath.CLOSEPOLY,
    ]
    return MplPath(verts, codes)


def draw_teaser_mpl(out_dir: Path) -> None:
    _mpl_setup()
    fig = plt.figure(figsize=(FIG_W_IN, FIG_H_IN), facecolor=WHITE)
    ax = fig.add_axes([0.015, 0.018, 0.97, 0.965])
    ax.set_xlim(0, CANVAS_W)
    ax.set_ylim(0, CANVAS_H)
    ax.set_aspect("equal", adjustable="box")
    ax.axis("off")

    ax.text(
        2,
        64.7,
        "CROSS-ENDOSCOPE NVS",
        ha="left",
        va="center",
        color=INK,
        fontsize=6.8,
        weight="bold",
    )
    source_lens = draw_mpl_endoscope(ax, 16.5, 56.4, -47, BLUE, "SOURCE ENDOSCOPE")
    target_lens = draw_mpl_endoscope(ax, 82.5, 56.4, -133, PURPLE, "TARGET ENDOSCOPE")

    ax.add_patch(
        FancyArrowPatch(
            (35, 58.5),
            (65, 58.5),
            arrowstyle="<->",
            mutation_scale=7,
            linestyle=(0, (3, 2)),
            linewidth=0.8,
            color=MUTED,
            zorder=3,
        )
    )
    ax.text(50, 60.2, "calibrated source → target", ha="center", va="bottom", color=MUTED, fontsize=6.2)

    ax.add_patch(
        Polygon(
            [source_lens, (16, 35.2), (60, 37.5)],
            closed=True,
            facecolor=BLUE_LIGHT,
            edgecolor=BLUE,
            linewidth=0.55,
            alpha=0.85,
            zorder=1,
        )
    )
    ax.add_patch(
        Polygon(
            [target_lens, (45, 37.4), (92, 35.0)],
            closed=True,
            facecolor=PURPLE_LIGHT,
            edgecolor=PURPLE,
            linewidth=0.55,
            alpha=0.82,
            zorder=1,
        )
    )
    ax.add_patch(PathPatch(tissue_path(), facecolor=TISSUE_LIGHT, edgecolor=TISSUE, lw=0.8, zorder=3))
    ax.plot([45, 53, 60], [37.6, 39.2, 38.1], color=TEAL, lw=2.3, solid_capstyle="round", zorder=4)
    ax.plot([65, 74, 84], [38.0, 39.8, 37.9], color=AMBER, lw=2.3, solid_capstyle="round", zorder=4)

    ray_x = [
        source_lens[0] + 0.42 * (48 - source_lens[0]),
        source_lens[0] + 0.56 * (48 - source_lens[0]),
        source_lens[0] + 0.70 * (48 - source_lens[0]),
    ]
    ray_y = [
        source_lens[1] + 0.42 * (38.2 - source_lens[1]),
        source_lens[1] + 0.56 * (38.2 - source_lens[1]),
        source_lens[1] + 0.70 * (38.2 - source_lens[1]),
    ]
    for i, (rx, ry) in enumerate(zip(ray_x, ray_y)):
        ax.add_patch(Circle((rx, ry), 0.55 + 0.10 * i, facecolor=BLUE, edgecolor=WHITE, lw=0.4, zorder=5))
    ax.text(31.5, 42.6, "depth distribution", ha="center", va="bottom", color=BLUE, fontsize=5.8)
    ax.text(52.2, 34.0, "source-supported", ha="center", va="top", color=TEAL, fontsize=5.9, weight="bold")
    ax.text(75.0, 34.0, "target-only", ha="center", va="top", color=AMBER, fontsize=5.9, weight="bold")

    rounded_box(ax, (39.0, 26.5), 22.0, 5.3, face=INK, edge=INK, radius=2.3, lw=0.7, z=7)
    ax.text(50, 29.15, "DistriNVS", ha="center", va="center", color=WHITE, fontsize=7.0, weight="bold", zorder=8)
    ax.add_patch(FancyArrowPatch((50, 34.0), (50, 31.9), arrowstyle="-|>", mutation_scale=7, lw=0.8, color=INK, zorder=6))

    rounded_box(ax, (2, 14.5), 63.0, 9.2, face=TEAL_LIGHT, edge=TEAL, radius=1.7, lw=0.75, z=2)
    ax.add_patch(Circle((7.0, 19.1), 2.2, facecolor=TEAL, edgecolor="none", zorder=4))
    ax.text(7.0, 19.1, "T", ha="center", va="center", color=WHITE, fontsize=7.2, weight="bold", zorder=5)
    ax.text(11.0, 20.3, "TRANSPORT OBSERVED", ha="left", va="center", color=TEAL, fontsize=6.5, weight="bold")
    ax.text(11.0, 17.5, "DSS: support  ·  variance  ·  collision", ha="left", va="center", color=INK, fontsize=5.8)
    rounded_box(ax, (49.6, 17.1), 12.7, 4.0, face=WHITE, edge=TEAL, radius=1.3, lw=0.6, z=4)
    ax.text(55.95, 19.05, "transport", ha="center", va="center", color=TEAL, fontsize=5.3, weight="bold", zorder=5)

    rounded_box(ax, (2, 3.1), 63.0, 9.2, face=AMBER_LIGHT, edge=AMBER, radius=1.7, lw=0.75, z=2)
    ax.add_patch(Circle((7.0, 7.7), 2.2, facecolor=AMBER, edgecolor="none", zorder=4))
    ax.text(7.0, 7.7, "S", ha="center", va="center", color=WHITE, fontsize=7.2, weight="bold", zorder=5)
    ax.text(11.0, 8.9, "SYNTHESIZE TARGET-ONLY", ha="left", va="center", color=AMBER, fontsize=6.5, weight="bold")
    ax.text(11.0, 6.1, "Fourier completion  ·  predicted risk", ha="left", va="center", color=INK, fontsize=5.8)
    rounded_box(ax, (52.2, 5.6), 10.1, 4.0, face=WHITE, edge=AMBER, radius=1.3, lw=0.6, z=4)
    ax.text(57.25, 7.6, "inferred", ha="center", va="center", color=AMBER, fontsize=5.3, weight="bold", zorder=5)

    ax.text(83.4, 25.0, "TARGET VIEW", ha="center", va="bottom", color=INK, fontsize=6.3, weight="bold")
    rounded_box(ax, (71.5, 3.1), 25.5, 20.6, face=WHITE, edge=INK, radius=2.0, lw=0.9, z=3)
    ax.add_patch(Wedge((83.8, 13.4), 8.1, 90, 270, facecolor=TEAL_LIGHT, edgecolor="none", zorder=4))
    ax.add_patch(Wedge((83.8, 13.4), 8.1, -90, 90, facecolor=AMBER_LIGHT, edgecolor="none", zorder=4))
    ax.add_patch(Circle((83.8, 13.4), 8.1, facecolor="none", edgecolor=RULE, lw=0.7, zorder=6))
    xx = [76.6, 79.0, 81.0, 83.0, 85.0, 88.0, 91.0]
    yy = [12.1, 14.0, 12.8, 15.2, 13.2, 14.8, 12.4]
    ax.plot(xx, yy, color=TISSUE, lw=2.3, solid_capstyle="round", zorder=7)
    ax.plot([83.8, 83.8], [5.8, 21.0], color=WHITE, lw=1.0, zorder=7)
    for j, h in enumerate([1.3, 2.1, 3.0]):
        ax.add_patch(Rectangle((87.1 + 1.25 * j, 7.6), 0.7, h, facecolor=AMBER, edgecolor="none", zorder=8))
    ax.text(78.6, 18.1, "transport", ha="center", va="center", color=TEAL, fontsize=4.5, weight="bold", zorder=8)
    ax.text(89.2, 18.1, "synthesis", ha="center", va="center", color=AMBER, fontsize=4.5, weight="bold", zorder=8)
    ax.text(89.0, 6.0, "risk", ha="center", va="center", color=AMBER, fontsize=5.5, weight="bold", zorder=8)

    for y, color in [(19.1, TEAL), (7.7, AMBER)]:
        ax.add_patch(FancyArrowPatch((65.4, y), (71.0, 13.4), arrowstyle="-|>", mutation_scale=7, lw=0.8, color=color, zorder=5))

    out_dir.mkdir(parents=True, exist_ok=True)
    common = dict(bbox_inches="tight", pad_inches=0.015, facecolor=WHITE)
    fig.savefig(out_dir / "distrinvs_teaser.svg", format="svg", **common)
    fig.savefig(out_dir / "distrinvs_teaser.pdf", format="pdf", **common)
    fig.savefig(out_dir / "distrinvs_teaser.png", format="png", dpi=600, **common)
    plt.close(fig)


def rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


SLIDE_W_IN = 10.0
SLIDE_H_IN = 6.67


def sx(x: float):
    return Inches(SLIDE_W_IN * x / CANVAS_W)


def sy(y: float):
    return Inches(SLIDE_H_IN * y / CANVAS_H)


def add_text(slide, x, y, w, h, text, *, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Arial", valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(sx(x), sy(y), sx(w), sy(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def style_shape(shape, fill, line, line_width=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def add_round_rect(slide, x, y, w, h, fill, line, line_width=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx(x), sy(y), sx(w), sy(h))
    return style_shape(shape, fill, line, line_width)


def add_line(slide, x1, y1, x2, y2, color, width=1.0):
    line = slide.shapes.add_connector(1, sx(x1), sy(y1), sx(x2), sy(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_ppt_endoscope(slide, x, y, angle, color, label):
    body = add_round_rect(slide, x - 6.5, y - 2.6, 12.5, 5.2, WHITE, color, 1.5)
    body.rotation = angle
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx(x - 3.3), sy(y - 1.1), sx(3.6), sy(2.2))
    style_shape(marker, color, color, 0.4)
    marker.rotation = angle
    theta = math.radians(angle)
    shaft_start = (x + 5.2 * math.cos(theta), y + 5.2 * math.sin(theta))
    shaft_end = (x + 13.4 * math.cos(theta), y + 13.4 * math.sin(theta))
    add_line(slide, shaft_start[0], shaft_start[1], shaft_end[0], shaft_end[1], color, 3.0)
    lens = slide.shapes.add_shape(MSO_SHAPE.OVAL, sx(shaft_end[0] - 1.0), sy(shaft_end[1] - 1.5), sx(2.0), sy(3.0))
    style_shape(lens, INK, WHITE, 0.6)
    lens.rotation = angle
    add_text(slide, x - 10, y - 8.2, 20, 3.2, label, size=11.5, color=color, bold=True, align=PP_ALIGN.CENTER)
    return shaft_end


def add_chevron(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, sx(x), sy(y), sx(w), sy(h))
    style_shape(shape, color, color, 0.4)
    return shape


def draw_teaser_pptx(out_dir: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(WHITE)

    add_text(slide, 2, 1.0, 60, 4.0, "CROSS-ENDOSCOPE NVS", size=13.5, color=INK, bold=True)
    # Editable FOV cones are inserted before devices so the cameras stay on top.
    fov_s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, sx(15), sy(18), sx(46), sy(22))
    style_shape(fov_s, BLUE_LIGHT, BLUE, 0.7)
    fov_t = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, sx(44), sy(18), sx(48), sy(22))
    style_shape(fov_t, PURPLE_LIGHT, PURPLE, 0.7)

    tissue = slide.shapes.add_shape(MSO_SHAPE.WAVE, sx(5), sy(35), sx(93), sy(7.0))
    style_shape(tissue, TISSUE_LIGHT, TISSUE, 1.1)
    source_support = slide.shapes.add_shape(MSO_SHAPE.ARC, sx(43), sy(34.7), sx(19), sy(5.0))
    source_support.fill.background()
    source_support.line.color.rgb = rgb(TEAL)
    source_support.line.width = Pt(3.0)
    target_only = slide.shapes.add_shape(MSO_SHAPE.ARC, sx(64), sy(34.6), sx(20), sy(5.2))
    target_only.fill.background()
    target_only.line.color.rgb = rgb(AMBER)
    target_only.line.width = Pt(3.0)

    add_ppt_endoscope(slide, 16.5, 13.0, 47, BLUE, "SOURCE ENDOSCOPE")
    add_ppt_endoscope(slide, 82.5, 13.0, 133, PURPLE, "TARGET ENDOSCOPE")
    add_line(slide, 36, 12.0, 64, 12.0, MUTED, 1.1)
    left_tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, sx(34.6), sy(10.9), sx(2.2), sy(2.2))
    style_shape(left_tri, MUTED, MUTED, 0.4)
    left_tri.rotation = 270
    right_tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, sx(63.2), sy(10.9), sx(2.2), sy(2.2))
    style_shape(right_tri, MUTED, MUTED, 0.4)
    right_tri.rotation = 90
    add_text(slide, 37, 7.6, 26, 3.2, "calibrated source → target", size=10.0, color=MUTED, align=PP_ALIGN.CENTER)

    add_text(slide, 42, 39.2, 22, 3.5, "source-supported", size=10.0, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 65, 39.2, 20, 3.5, "target-only", size=10.0, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    for i, (xx, yy) in enumerate([(29.0, 27.0), (34.0, 30.0), (39.0, 33.0)]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, sx(xx), sy(yy), sx(1.5 + 0.15 * i), sy(1.5 + 0.15 * i))
        style_shape(dot, BLUE, WHITE, 0.5)
    add_text(slide, 20, 30.6, 20, 3.2, "depth distribution", size=9.3, color=BLUE, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 39.0, 43.4, 22, 5.2, INK, INK, 0.6)
    add_text(slide, 40.0, 44.1, 20, 3.6, "DistriNVS", size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    down = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, sx(48.5), sy(40.8), sx(3.0), sy(2.4))
    style_shape(down, INK, INK, 0.4)
    down.rotation = 180

    add_round_rect(slide, 2, 50.1, 63, 7.1, TEAL_LIGHT, TEAL, 1.0)
    icon_t = slide.shapes.add_shape(MSO_SHAPE.OVAL, sx(4.4), sy(51.4), sx(4.4), sy(4.4))
    style_shape(icon_t, TEAL, TEAL, 0.5)
    add_text(slide, 4.4, 51.4, 4.4, 4.4, "T", size=12.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 10.2, 50.7, 38, 2.8, "TRANSPORT OBSERVED", size=10.7, color=TEAL, bold=True)
    add_text(slide, 10.2, 53.3, 40, 2.8, "DSS: support · variance · collision", size=9.2, color=INK)
    add_round_rect(slide, 52.0, 51.4, 10.5, 4.4, WHITE, TEAL, 0.8)
    add_text(slide, 52.5, 51.7, 9.5, 3.5, "transport", size=8.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 2, 58.6, 63, 7.1, AMBER_LIGHT, AMBER, 1.0)
    icon_s = slide.shapes.add_shape(MSO_SHAPE.OVAL, sx(4.4), sy(59.9), sx(4.4), sy(4.4))
    style_shape(icon_s, AMBER, AMBER, 0.5)
    add_text(slide, 4.4, 59.9, 4.4, 4.4, "S", size=12.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 10.2, 59.2, 40, 2.8, "SYNTHESIZE TARGET-ONLY", size=10.7, color=AMBER, bold=True)
    add_text(slide, 10.2, 61.8, 40, 2.8, "Fourier completion · predicted risk", size=9.2, color=INK)
    add_round_rect(slide, 52.0, 59.9, 10.5, 4.4, WHITE, AMBER, 0.8)
    add_text(slide, 53.0, 60.3, 8.5, 3.5, "inferred", size=8.3, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

    add_text(slide, 75.0, 47.0, 18.0, 3.0, "TARGET VIEW", size=10.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 71.5, 50.1, 25.5, 15.6, WHITE, INK, 1.1)
    add_round_rect(slide, 74.0, 52.0, 10.2, 11.8, TEAL_LIGHT, TEAL_LIGHT, 0.3)
    add_round_rect(slide, 83.7, 52.0, 10.8, 11.8, AMBER_LIGHT, AMBER_LIGHT, 0.3)
    add_text(slide, 74.4, 52.2, 9.2, 2.5, "transport", size=8.0, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 84.1, 52.2, 9.9, 2.5, "synthesis", size=8.0, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 75.5, 58.5, 92.8, 58.5, TISSUE, 3.0)
    for j, height in enumerate([1.4, 2.2, 3.0]):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx(87.0 + 1.4 * j), sy(60.4 - height), sx(0.75), sy(height))
        style_shape(bar, AMBER, AMBER, 0.3)
    add_text(slide, 86.5, 62.0, 6.0, 2.0, "risk", size=8.0, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_chevron(slide, 65.7, 53.1, 4.3, 3.0, TEAL)
    add_chevron(slide, 65.7, 61.6, 4.3, 3.0, AMBER)

    guide = prs.slides.add_slide(prs.slide_layouts[6])
    guide.background.fill.solid()
    guide.background.fill.fore_color.rgb = rgb(WHITE)
    add_text(guide, 5, 4, 90, 8, "DistriNVS teaser · editing guide", size=24, color=INK, bold=True)
    add_text(
        guide,
        5,
        13,
        88,
        10,
        "Recommended final size: 88 mm wide (single column). Keep source and target endoscopes physically separate; amber always denotes inferred target-only content.",
        size=15,
        color=MUTED,
    )
    palette = [(BLUE, "source"), (PURPLE, "target"), (TEAL, "transport / fusion"), (AMBER, "synthesis / risk"), (TISSUE, "tissue")]
    for i, (col, name) in enumerate(palette):
        xx = 6 + i * 18.5
        swatch = guide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx(xx), sy(28), sx(13.5), sy(8))
        style_shape(swatch, col, col, 0.5)
        add_text(guide, xx - 1, 37, 15.5, 5, name, size=12, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(guide, 5, 48, 90, 5, "Suggested caption", size=17, color=INK, bold=True)
    add_text(
        guide,
        5,
        54,
        90,
        11,
        "Cross-endoscope synthesis with explicit provenance. DSS transports source-supported anatomy and exposes visibility uncertainty; DistriNVS applies reliability-constrained soft fusion to transported and synthesized RGB-D.",
        size=14,
        color=MUTED,
    )

    out_dir.mkdir(parents=True, exist_ok=True)
    prs.save(out_dir / "distrinvs_teaser_editable.pptx")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    draw_teaser_mpl(args.output_dir)
    draw_teaser_pptx(args.output_dir)
    print(f"Wrote teaser assets to {args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
